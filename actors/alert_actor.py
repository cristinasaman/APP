from simgrid import Mailbox, this_actor, Engine
import openpyxl, os

class AlertActor:
    def __init__(self, alert_server: str):
        """
        alert_server: Name of the alert server.
        """
        
        self.alert_server = alert_server
        self.alert_mailbox = Mailbox.by_name(alert_server)
        self.flops_alert_log = 500e3
        
        self.total_log_computation_time = 0.0
        self.total_wait_time_for_alert = 0.0
        
        this_actor.info(f"AlertActor initialized. Listening on '{self.alert_server}'.")
        
    def __call__(self):
        this_actor.info(f"Started. Waiting for alerts on {self.alert_mailbox} at {Engine.clock:.3f}")
        
        while True:
            try:
                time_before_get = Engine.clock
                alert_data = self.alert_mailbox.get()
                self.total_wait_time_for_alert += (Engine.clock - time_before_get)
                
                this_actor.info(f"Received alert '{alert_data}' at {Engine.clock:.3f}")
                
                time_before_exec = Engine.clock
                this_actor.execute(self.flops_alert_log)
                self.total_log_computation_time += (Engine.clock - time_before_exec)
                
            except Exception as e:
                this_actor.error(f"Error while processing alert: {e}")
                break
        
        self._write_metrics_to_excel()
        this_actor.info(f"AlertActor stopping.")    

    def _write_metrics_to_excel(self):
        filename = "simulation_metrics.xlsx"
        sheet_name = "AlertMetrics"

        if not os.path.exists(filename):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
        else:
            wb = openpyxl.load_workbook(filename)
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(title=sheet_name)
                ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
            else:
                ws = wb[sheet_name]

        ws.append([
            f"{self.alert_server}", 
            round(self.total_log_computation_time , 6), 
            round(self.total_wait_time_for_alert , 6)
        ])

        wb.save(filename)                        