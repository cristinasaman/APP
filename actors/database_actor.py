from simgrid import Mailbox, this_actor, Engine 
import openpyxl, os

class DatabaseActor:
    def __init__(self, my_mailbox_name: str): 
        """
        my_mailbox_name: The mailbox name this DatabaseActor will listen on.
        """
        
        self.my_mailbox_name = my_mailbox_name
        try:
            self.mailbox = Mailbox.by_name(self.my_mailbox_name)
        except Exception as e:
            this_actor.error(f"Failed to get own mailbox '{self.my_mailbox_name}'. Actor cannot start.")
            raise e
        
        self.db_lookup_flops = 200e6  

        self.total_computation_time = 0.0
        self.total_wait_time_for_query = 0.0
        self.total_communication_send_time = 0.0 

        self.init_access_rules() 
        
        this_actor.info(f"DatabaseActor initialized. Listening on '{self.my_mailbox_name}'.")

    def __call__(self):
        this_actor.info(f"Started. Waiting for queries on '{self.my_mailbox_name}'.")
        while True:
            try:
                time_before_get = Engine.clock
                query_message = self.mailbox.get()

                if not isinstance(query_message, dict) or query_message.get("type") != "auth_query":
                    this_actor.warning(f"Received malformed or unexpected query type: {query_message}")
                    continue

                self.total_wait_time_for_query += (Engine.clock - time_before_get)

                face_id = query_message.get("face_id")
                zone_id = query_message.get("zone_id")
                reply_to_mailbox_name = query_message.get("reply_to_mailbox")

                if not all([face_id, zone_id, reply_to_mailbox_name]):
                    this_actor.error(f"Incomplete query received: {query_message}. Missing required fields.")
                    if reply_to_mailbox_name:
                        try:
                            error_response = {"type": "auth_result", "status": "Error_BadRequest", "queried_face_id": face_id, "queried_zone_id": zone_id}
                            Mailbox.by_name(reply_to_mailbox_name).put(error_response, 100)
                        except Exception as e_err_reply:
                            this_actor.error(f"Failed to send error reply for bad query to {reply_to_mailbox_name}: {e_err_reply}")
                    continue
                
                this_actor.info(f"Received query: FaceID='{face_id}', Zone='{zone_id}'. Reply to='{reply_to_mailbox_name}'.")
                
                time_before_exec = Engine.clock
                this_actor.execute(self.db_lookup_flops) 
                self.total_computation_time += (Engine.clock - time_before_exec)

                authorized = self.check_authorization(face_id, zone_id)
                
                response_status = "Authorized" if authorized else "Unauthorized"
                response_payload = {
                    "type": "auth_result",
                    "status": response_status,
                    "queried_face_id": face_id, 
                    "queried_zone_id": zone_id
                }
                
                try:
                    reply_mailbox = Mailbox.by_name(reply_to_mailbox_name)

                    response_message_size_bytes = 256 
                    
                    time_before_put = Engine.clock
                    reply_mailbox.put(response_payload, response_message_size_bytes) 
                    self.total_communication_send_time += (Engine.clock - time_before_put)
                    
                    this_actor.info(f"Sent response '{response_status}' for FaceID:'{face_id}', Zone:'{zone_id}' to '{reply_to_mailbox_name}'.")
                except Exception as e_reply:
                    this_actor.error(f"Failed to get reply mailbox '{reply_to_mailbox_name}' or send reply: {e_reply}")

            except Exception as e:
                this_actor.error(f"DatabaseActor error in main loop: {e}")
                break 
            except Exception as e_gen: 
                this_actor.error(f"DatabaseActor UNEXPECTED error in main loop: {e_gen}")
                
        self._write_metrics_to_excel()        
        this_actor.info("DatabaseActor stopping.")
        
    def init_access_rules(self):
        self.access_rules = {
            "zone_1": {
                "zone1_employee_": True, "zone1_visitor_": False, "zone1_manager_": True,
                "zone1_security_": True, "unknown_person_": False, "unknown_face": False,
                "low_confidence_face": False, "default": False
            },
            "zone_2": { 
                "zone2_employee_": True, "zone2_manager_": True, "zone1_security_": True, 
                "unknown_person_": False, "unknown_face": False, "low_confidence_face": False,
                "default": False
            },
            "zone_3": { 
                "zone3_employee_": True, "zone3_security_": True, "zone1_manager_": True, 
                "unknown_person_": False, "unknown_face": False, "low_confidence_face": False,
                "default": False
            },
            "default": {"default": False} 
        }
        
        self.special_users = {
            "blacklist": ["zone1_employee_042", "zone1_visitor_013"],
            "whitelist": ["zone1_manager_001", "zone1_security_007"] 
        }
        
        this_actor.info("Access rules initialized.")

    def check_authorization(self, face_id: str, zone_id: str) -> bool:
        if not face_id:
            this_actor.info(f"Authorization check for EMPTY FaceID in Zone '{zone_id}' -> DENIED (default)")
            return False
            
        if face_id in self.special_users["blacklist"]:
            this_actor.info(f"Access DENIED for '{face_id}' in Zone '{zone_id}' (blacklisted).")
            return False
        if face_id in self.special_users["whitelist"]:
            this_actor.info(f"Access GRANTED for '{face_id}' in Zone '{zone_id}' (whitelisted).")
            return True
        
        zone_specific_rules = self.access_rules.get(zone_id, self.access_rules["default"])
        
        for prefix, is_authorized in zone_specific_rules.items():
            if prefix != "default" and face_id.startswith(prefix):
                log_status = "GRANTED" if is_authorized else "DENIED"
                this_actor.info(f"Access {log_status} for '{face_id}' in Zone '{zone_id}' (rule: '{prefix}').")
                return is_authorized
                
        default_auth = zone_specific_rules.get("default", False)
        log_status = "GRANTED" if default_auth else "DENIED"
        this_actor.info(f"No specific rule for '{face_id}' in Zone '{zone_id}'. Applying zone default: {log_status}.")
        
        return default_auth        
    
    def _write_metrics_to_excel(self):
        filename = "simulation_metrics.xlsx"
        sheet_name = "DatabaseMetrics"

        if not os.path.exists(filename):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
        else:
            wb = openpyxl.load_workbook(filename)
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(title=sheet_name)
                ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
            else:
                ws = wb[sheet_name]

        ws.append([
            f"{self.my_mailbox_name}", 
            round(self.total_computation_time , 6), 
            round(self.total_wait_time_for_query , 6),
            round(self.total_communication_send_time , 6)
        ])

        wb.save(filename)                