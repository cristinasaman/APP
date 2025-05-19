from simgrid import Mailbox, this_actor, Engine
import openpyxl
import os

class ProcessingActorCPU:
    def __init__(self, 
                 name: str, 
                 my_mailbox_name: str, 
                 dispatcher_mailbox_name: str, 
                 database_mailbox_name: str, 
                 alert_mailbox_name: str,
                 od_accelerator_mb_name: str, fr_accelerator_mb_name: str):
        """
        name: Name of this CPU processing actor.
        my_mailbox_name: This actor's own mailbox name (to receive tasks from dispatcher, and replies).
        dispatcher_mailbox_name: Mailbox of the DispatcherActor (to send "ready" messages).
        database_mailbox_name: Mailbox of the DatabaseActor.
        alert_mailbox_name: Mailbox of the AlertActor.
        od_accelerator_mb_name: Mailbox of its dedicated OD accelerator.
        fr_accelerator_mb_name: Mailbox of its dedicated FR accelerator.
        """
        
        self.name = name
        self.my_mailbox_name = my_mailbox_name
        self.my_mailbox = Mailbox.by_name(my_mailbox_name)
        
        self.dispatcher_mailbox = Mailbox.by_name(dispatcher_mailbox_name)
        self.database_mailbox = Mailbox.by_name(database_mailbox_name)
        self.alert_mailbox = Mailbox.by_name(alert_mailbox_name)
        self.od_accelerator_mailbox = Mailbox.by_name(od_accelerator_mb_name)
        self.fr_accelerator_mailbox = Mailbox.by_name(fr_accelerator_mb_name)
        
        self.cpu_orchestration_cost = 1e7 
        self.db_query_flops = 200e6 
        self.alert_setup_flops = 2e6 
        self.control_message_size = 100 
        self.task_to_accel_metadata_size = 512 
        self.face_region_data_size_estimate = 50 * 1024 
        
        self.total_cpu_computation_time = 0.0   
        self.total_wait_time_for_task = 0.0     
        self.total_wait_time_on_od_reply = 0.0  
        self.total_wait_time_on_fr_reply = 0.0  
        self.total_wait_time_on_db_reply = 0.0  
        
        self.total_comm_time_to_od = 0.0        
        self.total_comm_time_to_fr = 0.0        
        self.total_comm_time_to_db = 0.0        
        self.total_comm_time_to_alert = 0.0     
        self.total_comm_time_to_dispatcher = 0.0
        
        this_actor.info(f"({self.name}) CPU Processor initialized with mailbox: '{my_mailbox_name}'.")


    def _signal_ready_to_dispatcher(self):
        ready_message = {
            "type": "worker_ready",
            "worker_mailbox": self.my_mailbox_name,
            "worker_name": self.name
        }
        this_actor.info(f"({self.name}) Signaling READY to dispatcher.")
        
        time_before_put = Engine.clock
        self.dispatcher_mailbox.put(ready_message, self.control_message_size)
        self.total_comm_time_to_dispatcher += (Engine.clock - time_before_put)
        
    def __call__(self):
        this_actor.info(f"({self.name}) CPU part started. Listening on '{self.my_mailbox_name}'.")
        self._signal_ready_to_dispatcher()

        while True:
            current_task_payload = None 
            try:
                time_before_get_task = Engine.clock
                current_task_payload = self.my_mailbox.get()
                if not isinstance(current_task_payload, dict) or current_task_payload.get("type") != "frame_task":
                    this_actor.warning(f"({self.name}) Received non-task message or malformed task: {current_task_payload}")
                    self._signal_ready_to_dispatcher() 
                    continue
                self.total_wait_time_for_task += (Engine.clock - time_before_get_task)
                
                time_before_exec = Engine.clock
                this_actor.execute(self.cpu_orchestration_cost)
                self.total_cpu_computation_time += (Engine.clock - time_before_exec)
                
                frame_content_sim = current_task_payload.get("data")
                camera_info = current_task_payload.get("camera_info", {})
                zone_id = camera_info.get("zone", "unknown_zone") 
                original_frame_size = current_task_payload.get("original_frame_size")

                this_actor.info(f"({self.name}) Received task: '{frame_content_sim}' from cam '{camera_info.get('id')}' in Z:{zone_id}.")

                od_task_msg = {
                    "type": "od_task",
                    "frame_payload": current_task_payload, #
                    "original_frame_size": original_frame_size,
                    "reply_to_mailbox": self.my_mailbox_name
                }
                this_actor.info(f"({self.name}) Offloading OD task for frame size {original_frame_size} to OD Accelerator.")

                time_before_put_od = Engine.clock
                self.od_accelerator_mailbox.put(od_task_msg, original_frame_size) 
                self.total_comm_time_to_od += (Engine.clock - time_before_put_od)
                
                time_before_get_od = Engine.clock
                od_response = self.my_mailbox.get()
                if not isinstance(od_response, dict) or od_response.get("type") != "od_result":
                    this_actor.error(f"({self.name}) Expected OD result, got: {od_response}. Aborting task.")
                    self._signal_ready_to_dispatcher()
                    continue
                self.total_wait_time_on_od_reply += (Engine.clock - time_before_get_od)
                
                time_before_exec = Engine.clock
                this_actor.execute(self.cpu_orchestration_cost) 
                self.total_cpu_computation_time += (Engine.clock - time_before_exec)
                od_result_data = od_response.get("result", {"person_detected": False})
                this_actor.info(f"({self.name}) Got OD result: Person detected = {od_result_data.get('person_detected')}")

                if od_result_data.get("person_detected"):
                    person_coordinates = od_result_data.get("coordinates")
                    
                    fr_task_msg = {
                        "type": "fr_task",
                        "frame_payload_info": current_task_payload, 
                        "face_region_size": self.face_region_data_size_estimate, 
                        "face_coordinates": person_coordinates,
                        "reply_to_mailbox": self.my_mailbox_name
                    }
                    this_actor.info(f"({self.name}) Offloading FR task (face region size {self.face_region_data_size_estimate}) to FR Accelerator.")

                    time_before_put_fr = Engine.clock
                    self.fr_accelerator_mailbox.put(fr_task_msg, self.face_region_data_size_estimate)
                    self.total_comm_time_to_fr += (Engine.clock - time_before_put_fr)

                    time_before_get_fr = Engine.clock
                    fr_response = self.my_mailbox.get() 
                    if not isinstance(fr_response, dict) or fr_response.get("type") != "fr_result":
                        this_actor.error(f"({self.name}) Expected FR result, got: {fr_response}. Aborting FR part.")

                        self._signal_ready_to_dispatcher()
                        continue
                    self.total_wait_time_on_fr_reply += (Engine.clock - time_before_get_fr)
                        
                    time_before_exec = Engine.clock
                    this_actor.execute(self.cpu_orchestration_cost) 
                    self.total_cpu_computation_time += (Engine.clock - time_before_exec)
                    face_id = fr_response.get("face_id", "unknown_face")
                    this_actor.info(f"({self.name}) Got FR result: Face ID = {face_id}")

                    if face_id != "unknown_face": 
                        db_query_payload = {
                            "type": "auth_query",
                            "face_id": face_id,
                            "zone_id": zone_id, 
                            "reply_to_mailbox": self.my_mailbox_name
                        }
                        this_actor.info(f"({self.name}) Querying DB for {face_id} in {zone_id}")
                        
                        time_before_put_db = Engine.clock
                        self.database_mailbox.put(db_query_payload, self.db_query_message_size) 
                        self.total_comm_time_to_db += (Engine.clock - time_before_put_db)
                        
                        time_before_get_db = Engine.clock
                        db_auth_response = self.my_mailbox.get() 
                        if not isinstance(db_auth_response, dict) or db_auth_response.get("type") != "auth_result":
                             this_actor.error(f"({self.name}) Expected DB auth result, got: {db_auth_response}")
                             self._signal_ready_to_dispatcher()
                             continue
                        self.total_wait_time_on_db_reply += (Engine.clock - time_before_get_db)
                        
                        time_before_exec = Engine.clock
                        this_actor.execute(self.cpu_orchestration_cost)
                        self.total_cpu_computation_time += (Engine.clock - time_before_exec)
                        auth_status = db_auth_response.get("status", "Error")
                        this_actor.info(f"({self.name}) DB Auth for {face_id} in {zone_id}: {auth_status}")

                        if auth_status == "Unauthorized":
                            self.send_alert(current_task_payload, face_id, zone_id, "Unauthorized access")
                    else: 
                        self.send_alert(current_task_payload, face_id, zone_id, "Unknown face detected")
                else: 
                    this_actor.info(f"({self.name}) No person detected by OD. Task complete.")

                self._signal_ready_to_dispatcher()

            except Exception as e:
                this_actor.error(f"({self.name}) SimgridError in main loop: {e}. Signaling ready.")
                try:
                    self._signal_ready_to_dispatcher()
                except Exception as sig_e:
                     this_actor.error(f"({self.name}) Failed to signal ready after error: {sig_e}")
                     break
                 
        self._write_metrics_to_excel()
            
        this_actor.info(f"({self.name}) CPU part stopping.")

    def send_alert(self, original_task_payload, face_id, zone_id, reason):
        camera_id = original_task_payload.get("camera_info", {}).get("id", "unknown_cam")
        alert_data = {
            "type": "alert",
            "timestamp": Engine.clock,
            "zone_id": zone_id,
            "camera_id": camera_id,
            "face_id": face_id,
            "reason": reason
        }
        this_actor.warning(f"({self.name}) SENDING ALERT: {alert_data}")
        try:
            time_before_exec = Engine.clock
            this_actor.execute(self.alert_setup_flops) 
            self.total_cpu_computation_time += (Engine.clock - time_before_exec)
            
            time_before_put = Engine.clock
            self.alert_mailbox.put(alert_data, self.alert_message_size)
            self.total_comm_time_to_alert += (Engine.clock - time_before_put)
        except Exception as e:
            this_actor.error(f"({self.name}) Failed to send alert: {e}")
            
    def _write_metrics_to_excel(self):
        filename = "simulation_metrics.xlsx"
        sheet_name = "ProcessorMetrics"

        if not os.path.exists(filename):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
        else:
            wb = openpyxl.load_workbook(filename)
            if sheet_name not in wb.sheetnames:
                ws = wb.create_sheet(title=sheet_name)
                ws.append(["Actor", "Computation Time (s)", "Communication Wait Time (s)"])
            else:
                ws = wb[sheet_name]

        ws.append([
            f"{self.name}", 
            round(self.total_cpu_computation_time , 6), 
            round(self.total_wait_time_for_task, 6),
            round(self.total_wait_time_on_od_reply, 6),
            round(self.total_wait_time_on_fr_reply, 6),
            round(self.total_wait_time_on_db_reply, 6),
            round(self.total_comm_time_to_od, 6),
            round(self.total_comm_time_to_fr, 6),
            round(self.total_comm_time_to_db, 6),
            round(self.total_comm_time_to_alert, 6),
            round(self.total_comm_time_to_dispatcher, 6)
        ])

        wb.save(filename)               